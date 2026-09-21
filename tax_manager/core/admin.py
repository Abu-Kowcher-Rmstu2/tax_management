from django.contrib import admin
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import (
    Taxpayer,
    TaxAssessment,
    TaxTransaction,
    TaxNotice,
    ArrearRegister,
    NoticeSetting
)

# ---------------------------------------------------------
# Reusable Excel Export Helper
# ---------------------------------------------------------
def create_excel_response(filename, headers, rows_data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    header_fill = PatternFill(start_color="006A4E", end_color="006A4E", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    data_font = Font(name="Arial", size=10)
    for row in rows_data:
        ws.append(row)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = data_font

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response


# ---------------------------------------------------------
# Export Actions
# ---------------------------------------------------------
@admin.action(description="Export Selected to Excel (এক্সেল ডাউনলোড)")
def export_taxpayers_to_excel(modeladmin, request, queryset):
    headers = [
        "TIN", "Name", "Phone", "Email", "Holding/Circle",
        "Address Line 1", "Address Line 2", "Address Line 3", "Active Status"
    ]
    rows = []
    for item in queryset:
        rows.append([
            item.tin,
            item.name,
            item.phone,
            item.email,
            item.holding_or_circle,
            item.address_line1,
            item.address_line2,
            item.address_line3,
            "Active" if item.is_active else "Inactive"
        ])
    return create_excel_response("Taxpayers_Export", headers, rows)


@admin.action(description="Export Selected to Excel (এক্সেল ডাউনলোড)")
def export_transactions_to_excel(modeladmin, request, queryset):
    headers = [
        "Challan No", "TIN", "Taxpayer Name", "Transaction Type",
        "Fiscal Year", "Amount (Tk)", "Payment Mode", "Bank & Branch",
        "Challan Date", "Payment Date", "Remarks"
    ]
    rows = []
    for item in queryset.select_related('taxpayer'):
        rows.append([
            item.challan_no,
            item.taxpayer.tin,
            item.taxpayer.name,
            item.get_transaction_type_display(),
            item.fiscal_year,
            float(item.amount),
            item.get_payment_mode_display(),
            item.bank_and_branch,
            str(item.challan_date) if item.challan_date else '',
            str(item.payment_date) if item.payment_date else '',
            item.remarks
        ])
    return create_excel_response("Tax_Transactions_Export", headers, rows)


@admin.action(description="Export Selected to Excel (এক্সেল ডাউনলোড)")
def export_assessments_to_excel(modeladmin, request, queryset):
    headers = [
        "TIN", "Taxpayer Name", "Fiscal Year",
        "Demanded Amount (Tk)", "Arrear Brought Forward (Tk)"
    ]
    rows = []
    for item in queryset.select_related('taxpayer'):
        rows.append([
            item.taxpayer.tin,
            item.taxpayer.name,
            item.fiscal_year,
            float(item.demanded_amount),
            float(item.arrear_brought_forward)
        ])
    return create_excel_response("Tax_Assessments_Export", headers, rows)


@admin.action(description="Export Selected to Excel (বকেয়া রেজিস্টার এক্সেল ডাউনলোড)")
def export_arrears_to_excel(modeladmin, request, queryset):
    headers = [
        "TIN",
        "Taxpayer Name",
        "Assessment Year",
        "Arrear Type",
        "Disposal Section",
        "Disposal Date",
        "Demand (Tk)",
        "Fine (Tk)",
        "Total Arrear (Tk)",
        "Collection (Tk)",
        "Due Amount (Tk)",
        "Created At",
    ]
    rows = []
    for item in queryset.select_related('taxpayer'):
        # Format disposal date cleanly
        disposal_date_str = item.disposal_date.strftime("%d-%m-%Y") if item.disposal_date else ""
        created_at_str = item.created_at.strftime("%d-%m-%Y %H:%M") if item.created_at else ""

        rows.append([
            item.taxpayer.tin,
            item.taxpayer.name,
            item.assessment_year,
            item.get_arrear_type_display(),
            item.disposal_section or "",
            disposal_date_str,
            float(item.demand or 0.0),
            float(item.fine or 0.0),
            float(item.total_arrear or 0.0),
            float(item.collection or 0.0),
            float(item.due_amount or 0.0),
            created_at_str,
        ])
    return create_excel_response("Arrear_Register_Export", headers, rows)


@admin.register(ArrearRegister)
class ArrearRegisterAdmin(admin.ModelAdmin):
    list_display = (
        'get_tin',
        'get_taxpayer_name',
        'assessment_year',
        'arrear_type',
        'disposal_section',
        'disposal_date',
        'demand',
        'fine',
        'total_arrear',
        'collection',
        'due_amount',
        'created_at',
    )
    list_filter = ('arrear_type', 'assessment_year', 'disposal_date', 'created_at')
    search_fields = ('taxpayer__tin', 'taxpayer__name', 'assessment_year', 'disposal_section')
    readonly_fields = ('total_arrear', 'due_amount', 'created_at', 'updated_at')
    autocomplete_fields = ['taxpayer']
    ordering = ('-created_at',)
    actions = [export_arrears_to_excel]

    @admin.display(description='টিআইএন (TIN)', ordering='taxpayer__tin')
    def get_tin(self, obj):
        return obj.taxpayer.tin

    @admin.display(description='করদাতার নাম', ordering='taxpayer__name')
    def get_taxpayer_name(self, obj):
        return obj.taxpayer.name

# ---------------------------------------------------------
# ModelAdmin Registrations
# ---------------------------------------------------------
@admin.register(Taxpayer)
class TaxpayerAdmin(admin.ModelAdmin):
    list_display = ('tin', 'name', 'phone', 'holding_or_circle', 'address_line2', 'address_line3', 'is_active')
    search_fields = ('tin', 'name', 'phone', 'holding_or_circle', 'address_line1', 'address_line2', 'address_line3')
    list_filter = ('is_active', 'holding_or_circle')
    actions = [export_taxpayers_to_excel]


@admin.register(TaxTransaction)
class TaxTransactionAdmin(admin.ModelAdmin):
    list_display = ('challan_no', 'taxpayer', 'transaction_type', 'amount', 'challan_date', 'payment_date', 'payment_mode')
    list_filter = ('transaction_type', 'payment_mode', 'fiscal_year', 'payment_date')
    search_fields = ('challan_no', 'taxpayer__name', 'taxpayer__tin')
    date_hierarchy = 'payment_date'
    actions = [export_transactions_to_excel]


@admin.register(TaxAssessment)
class TaxAssessmentAdmin(admin.ModelAdmin):
    list_display = ('taxpayer', 'fiscal_year', 'demanded_amount', 'arrear_brought_forward')
    list_filter = ('fiscal_year',)
    search_fields = ('taxpayer__tin', 'taxpayer__name')
    actions = [export_assessments_to_excel]




@admin.register(TaxNotice)
class TaxNoticeAdmin(admin.ModelAdmin):
    list_display = ('reference_no', 'taxpayer', 'notice_type', 'arrear_amount', 'issue_date', 'due_date', 'is_served')
    list_filter = ('notice_type', 'is_served')
    search_fields = ('reference_no', 'taxpayer__tin', 'taxpayer__name')


@admin.register(NoticeSetting)
class NoticeSettingAdmin(admin.ModelAdmin):
    list_display = ('officer_name', 'designation', 'sender_circle', 'phone_number', 'is_active')
    list_editable = ('is_active',)
